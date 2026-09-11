"""
Exporta Markdown de app_desktop/docs/ a Word (.docx) en docs/word/.

Uso:
    cd app_desktop/docs
    pip install python-docx
    python _tools/export_to_docx.py

No forma parte del runtime de portería (no añadir a requirements.txt de la app).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:
    print(
        "Falta python-docx. Instale con: pip install python-docx",
        file=sys.stderr,
    )
    sys.exit(1)

DOCS_ROOT = Path(__file__).resolve().parent.parent
WORD_ROOT = DOCS_ROOT / "word"
SKIP_DIRS = {"_tools", "word", "mockups"}


def _set_run_font(run, *, bold: bool = False, mono: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(10 if mono else 11)
    run.font.name = "Consolas" if mono else "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), run.font.name)
    rFonts.set(qn("w:hAnsi"), run.font.name)


def _add_heading(doc: Document, text: str, level: int) -> None:
    level = max(1, min(level, 3))
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _set_run_font(run, bold=True)


def _add_para(doc: Document, text: str, *, mono: bool = False) -> None:
    p = doc.add_paragraph()
    # Inline code `...`
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = p.add_run(part[1:-1])
            _set_run_font(run, mono=True)
        else:
            # Bold **...**
            sub = re.split(r"(\*\*[^*]+\*\*)", part)
            for s in sub:
                if s.startswith("**") and s.endswith("**") and len(s) >= 4:
                    run = p.add_run(s[2:-2])
                    _set_run_font(run, bold=True)
                else:
                    run = p.add_run(s)
                    _set_run_font(run, mono=mono)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell_text = row[j] if j < len(row) else ""
            cell_text = re.sub(r"\*+", "", cell_text).strip()
            cell_text = re.sub(r"`([^`]+)`", r"\1", cell_text)
            table.rows[i].cells[j].text = cell_text
            for para in table.rows[i].cells[j].paragraphs:
                for run in para.runs:
                    _set_run_font(run, bold=(i == 0))


def _parse_table_block(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        cells = [c.strip() for c in line.strip("|").split("|")]
        # Skip separator |---|---|
        if all(re.match(r"^:?-+:?$", c) for c in cells if c):
            i += 1
            continue
        rows.append(cells)
        i += 1
    return rows, i


def md_to_docx(md_path: Path, out_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    in_code = False
    code_lang = ""
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line.strip()[3:].strip()
                code_buf = []
            else:
                in_code = False
                label = f"[{code_lang}]" if code_lang else "[código]"
                _add_para(doc, label, mono=True)
                for cl in code_buf:
                    _add_para(doc, cl if cl else " ", mono=True)
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines):
            rows, next_i = _parse_table_block(lines, i)
            if rows:
                _add_table(doc, rows)
                i = next_i
                continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            _add_heading(doc, m.group(2).strip(), level)
            i += 1
            continue

        if re.match(r"^[-*]\s+", line.strip()):
            item = re.sub(r"^[-*]\s+", "", line.strip())
            p = doc.add_paragraph(style="List Bullet")
            # reuse inline parsing
            parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", item)
            for part in parts:
                if part.startswith("`") and part.endswith("`"):
                    run = p.add_run(part[1:-1])
                    _set_run_font(run, mono=True)
                elif part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    _set_run_font(run, bold=True)
                else:
                    run = p.add_run(part)
                    _set_run_font(run)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", line.strip()):
            item = re.sub(r"^\d+\.\s+", "", line.strip())
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(re.sub(r"[*`]", "", item))
            _set_run_font(run)
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.strip() == "":
            i += 1
            continue

        _add_para(doc, line.strip())
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def collect_md_files() -> list[Path]:
    files: list[Path] = []
    for path in sorted(DOCS_ROOT.rglob("*.md")):
        rel_parts = path.relative_to(DOCS_ROOT).parts
        if rel_parts[0] in SKIP_DIRS:
            continue
        files.append(path)
    return files


def main() -> int:
    files = collect_md_files()
    if not files:
        print("No se encontraron .md para exportar.", file=sys.stderr)
        return 1

    WORD_ROOT.mkdir(parents=True, exist_ok=True)
    print(f"Exportando {len(files)} archivos -> {WORD_ROOT}")
    for md in files:
        rel = md.relative_to(DOCS_ROOT)
        out = WORD_ROOT / rel.with_suffix(".docx")
        md_to_docx(md, out)
        print(f"  OK {rel} -> word/{rel.with_suffix('.docx')}")
    print("Listo.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
